from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.units import EMU_to_pixels
from PIL import Image

from skin_cancer.webapp.extensions import db
from skin_cancer.webapp.models import Analysis, AnalysisResult, CaseImage, UploadBatch

EXCEL_URL = "/api/upload-batches/{batch_id}/analysis.xlsx"
EXPECTED_HEADERS = [
    "ID",
    "ẢNH GỐC",
    "GRAD-CAM",
    "TÊN FILE",
    "TRẠNG THÁI",
    "DỰ ĐOÁN",
    "TIN CẬY",
    "PHÂN LOẠI",
    "XỬ LÝ (ms)",
]


def _write_image(path, size, color):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)


def _workbook(response):
    return load_workbook(BytesIO(response.data))


def test_analysis_excel_exports_styled_partial_batch_with_embedded_images(
    app, client, users, case_graph, login
):
    record = case_graph["own"]
    upload_root = Path(app.config["UPLOAD_FOLDER"])
    gradcam_root = upload_root.parent / "gradcam"
    original_path = upload_root / "images" / "lesion-wide.png"
    gradcam_path = gradcam_root / "analysis.png"
    _write_image(original_path, (200, 100), "#AA3333")
    _write_image(gradcam_path, (100, 200), "#3344AA")
    app.config["GRADCAM_FOLDER"] = str(gradcam_root)

    with app.app_context():
        batch = db.session.get(UploadBatch, record["batch_id"])
        image = db.session.get(CaseImage, record["image_id"])
        image.image_path = str(original_path)
        image.original_filename = "lesion-wide.png"
        image.status = "COMPLETED"

        analysis = db.session.get(Analysis, record["analysis_id"])
        analysis.status = "COMPLETED"
        analysis.gradcam_path = str(gradcam_path)
        db.session.add(
            AnalysisResult(
                analysis_id=analysis.id,
                model_name="EfficientNet-B0",
                binary_result=1,
                binary_probability=0.91,
                predicted_class="AK",
                predicted_probability=0.973708,
                class_probabilities={"AK": 0.973708},
                inference_time_ms=236.66,
            )
        )

        pending = CaseImage(
            case_id=record["case_id"],
            batch_id=batch.id,
            image_path=str(upload_root / "images" / "missing.png"),
            original_filename="pending.png",
            mime_type="image/png",
            file_size=0,
            uploaded_by=users["doctor"],
            status="UPLOADED",
        )
        db.session.add(pending)
        batch.total_images = 2
        batch.success_count = 2
        db.session.commit()
        pending_id = pending.id

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))

    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "batch_1_analysis.xlsx" in response.headers["Content-Disposition"]

    workbook = _workbook(response)
    assert workbook.sheetnames == ["ẢNH TRONG BATCH"]
    sheet = workbook["ẢNH TRONG BATCH"]
    assert sheet["A1"].value == "DERMASCAN AI — BATCH ANALYSIS"
    assert sheet["A2"].value == "BATCH #1 — IMAGE ANALYSIS REPORT"
    assert [sheet.cell(5, column).value for column in range(1, 10)] == EXPECTED_HEADERS
    assert sheet.freeze_panes == "A6"
    assert sheet["A1"].fill.fgColor.rgb.endswith("2BA149")
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.orientation == "landscape"

    assert sheet["A6"].value == record["image_id"]
    assert sheet["D6"].value == "lesion-wide.png"
    assert sheet["E6"].value == "COMPLETED"
    assert sheet["F6"].value == "AK"
    assert sheet["G6"].value == 0.973708
    assert sheet["G6"].number_format == "0.00%"
    assert sheet["H6"].value == "MALIGNANT"
    assert sheet["I6"].value == 236.66

    assert sheet["A7"].value == pending_id
    assert sheet["C7"].value == "Chưa phân tích"
    assert sheet["E7"].value == "UPLOADED"
    assert sheet["F7"].value is None
    assert sheet["G7"].value is None
    assert sheet["H7"].value == "Chưa phân tích"
    assert sheet["I7"].value is None

    assert len(sheet._images) == 2
    embedded_by_column = {image.anchor._from.col: image for image in sheet._images}
    original = embedded_by_column[1]
    gradcam = embedded_by_column[2]
    original_width = EMU_to_pixels(original.anchor.ext.cx)
    original_height = EMU_to_pixels(original.anchor.ext.cy)
    gradcam_width = EMU_to_pixels(gradcam.anchor.ext.cx)
    gradcam_height = EMU_to_pixels(gradcam.anchor.ext.cy)
    assert original_width / original_height == 2
    assert gradcam_height / gradcam_width == 2
    assert max(original_width, original_height) <= 96
    assert max(gradcam_width, gradcam_height) <= 96


def test_analysis_excel_handles_missing_or_corrupt_images_without_failing(
    app, client, case_graph, login
):
    record = case_graph["own"]
    upload_root = Path(app.config["UPLOAD_FOLDER"])
    gradcam_root = upload_root.parent / "gradcam"
    corrupt_original = upload_root / "images" / "corrupt.png"
    corrupt_gradcam = gradcam_root / "corrupt.png"
    corrupt_original.parent.mkdir(parents=True, exist_ok=True)
    corrupt_gradcam.parent.mkdir(parents=True, exist_ok=True)
    corrupt_original.write_text("not an image", encoding="utf-8")
    corrupt_gradcam.write_text("not an image", encoding="utf-8")
    app.config["GRADCAM_FOLDER"] = str(gradcam_root)

    with app.app_context():
        image = db.session.get(CaseImage, record["image_id"])
        image.image_path = str(corrupt_original)
        analysis = db.session.get(Analysis, record["analysis_id"])
        analysis.status = "COMPLETED"
        analysis.gradcam_path = str(corrupt_gradcam)
        db.session.commit()

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))

    assert response.status_code == 200
    sheet = _workbook(response).active
    assert sheet["B6"].value == "Không thể đọc ảnh gốc"
    assert sheet["C6"].value == "Không thể đọc Grad-CAM"
    assert sheet._images == []
    assert str(corrupt_original) not in response.data.decode("latin-1")


def test_analysis_excel_returns_404_for_unknown_batch(client, users, login):
    login("doctor")
    assert client.get(EXCEL_URL.format(batch_id=999999)).status_code == 404


def test_excel_dependencies_are_declared_for_database_runtime():
    requirements = Path("requirements-database.txt").read_text(encoding="utf-8")
    assert "openpyxl" in requirements.lower()
    assert "pillow" in requirements.lower()


def test_analysis_excel_neutralizes_spreadsheet_formula_filename(
    app, client, case_graph, login
):
    record = case_graph["own"]
    with app.app_context():
        image = db.session.get(CaseImage, record["image_id"])
        image.original_filename = '=HYPERLINK("https://example.invalid")'
        db.session.commit()

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))
    filename_cell = _workbook(response).active["D6"]

    assert filename_cell.value.startswith("'=")
    assert filename_cell.data_type == "s"


def test_analysis_excel_normalizes_webp_to_small_png_thumbnail(
    app, client, case_graph, login
):
    record = case_graph["own"]
    upload_root = Path(app.config["UPLOAD_FOLDER"])
    webp_path = upload_root / "images" / "large-lesion.webp"
    webp_path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (1200, 600), "#AA3333").save(webp_path, format="WEBP")

    with app.app_context():
        image = db.session.get(CaseImage, record["image_id"])
        image.image_path = str(webp_path)
        image.original_filename = webp_path.name
        db.session.commit()

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))

    assert response.status_code == 200
    with ZipFile(BytesIO(response.data)) as archive:
        media_names = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(media_names) == 1
        assert media_names[0].endswith(".png")
        with Image.open(BytesIO(archive.read(media_names[0]))) as embedded:
            embedded.load()
            assert embedded.format == "PNG"
            assert embedded.size == (96, 48)


def test_analysis_excel_rejects_truncated_jpeg_before_embedding(
    app, client, case_graph, login
):
    record = case_graph["own"]
    upload_root = Path(app.config["UPLOAD_FOLDER"])
    jpeg_path = upload_root / "images" / "truncated.jpg"
    jpeg_path.parent.mkdir(parents=True, exist_ok=True)
    buffer = BytesIO()
    Image.new("RGB", (1200, 600), "#AA3333").save(buffer, format="JPEG")
    jpeg_path.write_bytes(buffer.getvalue()[:-300])

    with app.app_context():
        image = db.session.get(CaseImage, record["image_id"])
        image.image_path = str(jpeg_path)
        image.original_filename = jpeg_path.name
        db.session.commit()

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))

    assert response.status_code == 200
    sheet = _workbook(response).active
    assert sheet["B6"].value == "Không thể đọc ảnh gốc"
    assert sheet._images == []


def test_analysis_excel_removes_xml_control_characters_from_cells(
    app, client, case_graph, login
):
    record = case_graph["own"]
    with app.app_context():
        image = db.session.get(CaseImage, record["image_id"])
        image.original_filename = "lesion\x01scan.png"
        db.session.commit()

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))

    assert response.status_code == 200
    assert _workbook(response).active["D6"].value == "lesionscan.png"


def test_analysis_excel_neutralizes_formula_after_xml_control_character(
    app, client, case_graph, login
):
    record = case_graph["own"]
    with app.app_context():
        image = db.session.get(CaseImage, record["image_id"])
        image.original_filename = '\x01=HYPERLINK("https://example.invalid")'
        db.session.commit()

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))
    filename_cell = _workbook(response).active["D6"]

    assert filename_cell.value.startswith("'=")
    assert filename_cell.data_type == "s"


def test_analysis_excel_embeds_valid_cmyk_jpeg(
    app, client, case_graph, login
):
    record = case_graph["own"]
    upload_root = Path(app.config["UPLOAD_FOLDER"])
    jpeg_path = upload_root / "images" / "cmyk-lesion.jpg"
    jpeg_path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("CMYK", (120, 60), (0, 80, 80, 20)).save(jpeg_path, format="JPEG")

    with app.app_context():
        image = db.session.get(CaseImage, record["image_id"])
        image.image_path = str(jpeg_path)
        image.original_filename = jpeg_path.name
        db.session.commit()

    login("doctor")
    response = client.get(EXCEL_URL.format(batch_id=record["batch_id"]))

    assert response.status_code == 200
    sheet = _workbook(response).active
    assert sheet["B6"].value is None
    assert len(sheet._images) == 1
